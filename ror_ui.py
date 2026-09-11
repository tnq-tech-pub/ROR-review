import csv
import fcntl
import io
import json
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook

import ror_implementor


def _env_path(name: str, default: str) -> Path:
    return Path(os.getenv(name, default)).expanduser()

OUTPUT_ROOT = _env_path("OUTPUT_ROOT", "ror_results")
BATCHES_ROOT = _env_path("BATCHES_ROOT", str(OUTPUT_ROOT / "batches"))
CSV_DIR = _env_path("CSV_DIR", str(OUTPUT_ROOT / "verifications"))
INSERTIONS_CSV = Path(
    os.getenv("INSERTIONS_CSV", str(OUTPUT_ROOT / "ror_insertions.csv"))
).expanduser()
REMARKS_CSV = Path(
    os.getenv("REMARKS_CSV", str(OUTPUT_ROOT / "ror_remarks.csv"))
).expanduser()

INSERTION_RECORD_TYPE = "ror_org_insertion_v1"
INSERTION_REQUIRED_KEYS = (
    "organization_name",
    "organization_website",
    "organization_domain",
    "organization_type",
    "city_primary",
    "country_primary",
)
INSERTION_OPTIONAL_KEYS = (
    "names_other_languages",
    "name_variations",
    "acronym",
    "wikipedia_page",
    "wikidata_id",
    "isni_id",
    "grid_id",
    "crossref_funder_id",
    "year_established",
    "parent_org_ror",
    "child_org_ror",
    "related_org_ror",
    "requestor_comments",
)
INSERTIONS_TEMPLATE_XLSX = Path(
    os.getenv(
        "INSERTIONS_TEMPLATE_XLSX",
        str(
            Path(__file__).resolve().parent
            / "PUBLIC ROR Bulk Processing Template - New Records.xlsx"
        ),
    )
).expanduser()
INSERTIONS_DOWNLOAD_NAME = "PUBLIC ROR Bulk Processing Template - New Records.xlsx"
INSERTIONS_TEMPLATE_SHEET = "Data provided by requestor"
INSERTIONS_TEMPLATE_FIRST_DATA_ROW = 4
INSERTION_TEMPLATE_COLS = {
    "organization_name": 2,
    "names_other_languages": 3,
    "name_variations": 4,
    "acronym": 5,
    "organization_website": 6,
    "organization_domain": 7,
    "wikipedia_page": 9,
    "wikidata_id": 10,
    "isni_id": 11,
    "grid_id": 12,
    "crossref_funder_id": 13,
    "organization_type": 14,
    "year_established": 15,
    "parent_org_ror": 16,
    "child_org_ror": 17,
    "related_org_ror": 18,
    "city_primary": 19,
    "country_primary": 20,
    "requestor_comments": 21,
}

app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True
if os.getenv("FLASK_STATIC_NO_CACHE", "").lower() in ("1", "true", "yes"):
    app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


def _json_load(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _batch_result_files(batch_id: Optional[str] = None) -> List[Path]:
    if not BATCHES_ROOT.exists():
        return []
    if batch_id:
        batch_dir = BATCHES_ROOT / batch_id
        if not batch_dir.is_dir():
            return []
        return sorted(batch_dir.glob("*.json"))

    files: List[Path] = []
    for batch_dir in sorted(BATCHES_ROOT.iterdir()):
        if batch_dir.is_dir():
            files.extend(sorted(batch_dir.glob("*.json")))
    return files


def _result_detail_path(batch_id: str, affiliation_id: str) -> Path:
    return BATCHES_ROOT / batch_id / f"{affiliation_id}.json"


def _extract_items(response: Dict[str, Any]) -> List[Dict[str, Any]]:
    items = response.get("items")
    if isinstance(items, list):
        return items
    return []


def _top_affiliation_item(response: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    items = _extract_items(response)
    if not items:
        return None
    for item in items:
        if item.get("chosen") is True:
            return item
    return items[0]


def _first_item_score_and_substring(response: Dict[str, Any]) -> tuple[str, str]:
    item = _top_affiliation_item(response)
    if not item:
        return "", ""
    score = item.get("score")
    if score is None:
        score_str = ""
    elif isinstance(score, (int, float)):
        score_str = str(score)
    else:
        score_str = str(score)
    sub = item.get("substring")
    if isinstance(sub, str):
        sub_str = sub
    elif sub is None:
        sub_str = ""
    else:
        sub_str = str(sub)
    return score_str, sub_str


def _first_item_matching_type(response: Dict[str, Any]) -> str:
    item = _top_affiliation_item(response)
    if not item:
        return "NO_RESULT"
    matching_type = item.get("matching_type")
    if matching_type is None:
        return ""
    return str(matching_type).strip()


def _chosen_any(response: Dict[str, Any]) -> bool:
    for item in _extract_items(response):
        if item.get("chosen") is True:
            return True
    return False


def _verification_csv_path(dt: Optional[datetime] = None) -> Path:
    use_dt = dt or datetime.now()
    date_part = use_dt.strftime("%Y-%m-%d")
    return CSV_DIR / f"ror_verifications_{date_part}.csv"


def _load_verifications() -> Dict[str, Dict[str, Any]]:
    csv_path = _verification_csv_path()
    if not csv_path.exists():
        return {}
    records: Dict[str, Dict[str, Any]] = {}
    with csv_path.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            records[row["row_id"]] = row
    return records


def _load_insertions_latest() -> Dict[str, str]:
    if not INSERTIONS_CSV.exists():
        return {}
    latest: Dict[str, str] = {}
    with INSERTIONS_CSV.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            rid = row.get("row_id") or ""
            if rid:
                latest[rid] = (row.get("insertion_text") or "").strip()
    return latest


def _is_csv_header_data_row(row: Dict[str, str]) -> bool:
    return (row.get("row_id") or "").strip().lower() == "row_id"


def _append_csv_row(csv_path: Path, fieldnames: List[str], row: Dict[str, Any]) -> None:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("a+", newline="", encoding="utf-8") as handle:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
        try:
            handle.seek(0, os.SEEK_END)
            write_header = handle.tell() == 0
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            if write_header:
                writer.writeheader()
            writer.writerow(row)
        finally:
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


def _load_remarks_latest() -> Dict[str, str]:
    if not REMARKS_CSV.exists():
        return {}
    latest: Dict[str, str] = {}
    with REMARKS_CSV.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if _is_csv_header_data_row(row):
                continue
            rid = row.get("row_id") or ""
            if rid:
                latest[rid] = (row.get("remarks") or "").strip()
    return latest


def _append_remarks(row: Dict[str, Any]) -> None:
    fieldnames = ["row_id", "timestamp", "batch_id", "affiliation_id", "remarks"]
    _append_csv_row(REMARKS_CSV, fieldnames, row)


def _normalize_insertion_record(raw: Dict[str, Any]) -> tuple[Optional[Dict[str, str]], Optional[str]]:
    out: Dict[str, str] = {"_type": INSERTION_RECORD_TYPE}
    for key in INSERTION_REQUIRED_KEYS + INSERTION_OPTIONAL_KEYS:
        val = raw.get(key)
        if val is None:
            out[key] = ""
        elif isinstance(val, (dict, list)):
            return None, f"Invalid value for field: {key}"
        else:
            out[key] = str(val).strip()
    for key in INSERTION_REQUIRED_KEYS:
        if not out[key]:
            return None, f"Required field is empty: {key}"
    year = out.get("year_established", "")
    if year and (len(year) != 4 or not year.isdigit()):
        return None, "Year established must be four digits (YYYY) or left empty."
    return out, None


def _parse_insertion_text(text: str) -> Optional[Dict[str, str]]:
    if not text:
        return None
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return None
    if not isinstance(parsed, dict) or parsed.get("_type") != INSERTION_RECORD_TYPE:
        return None
    return parsed


def _load_insertion_records_list() -> List[Dict[str, str]]:
    if not INSERTIONS_CSV.exists():
        return []
    latest: Dict[str, Dict[str, str]] = {}
    with INSERTIONS_CSV.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            rid = row.get("row_id") or ""
            rec = _parse_insertion_text((row.get("insertion_text") or "").strip())
            if rid and rec:
                latest[rid] = rec
    return list(latest.values())


def _build_insertions_xlsx(records: List[Dict[str, str]]) -> io.BytesIO:
    if not INSERTIONS_TEMPLATE_XLSX.exists():
        raise FileNotFoundError(f"Insertion template not found: {INSERTIONS_TEMPLATE_XLSX}")
    workbook = load_workbook(INSERTIONS_TEMPLATE_XLSX)
    worksheet = workbook[INSERTIONS_TEMPLATE_SHEET]
    row_idx = INSERTIONS_TEMPLATE_FIRST_DATA_ROW
    for record in records:
        for field, col in INSERTION_TEMPLATE_COLS.items():
            worksheet.cell(row=row_idx, column=col, value=record.get(field, "") or "")
        row_idx += 1
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _append_insertion(row: Dict[str, Any]) -> None:
    fieldnames = [
        "row_id",
        "timestamp",
        "batch_id",
        "affiliation_id",
        "insertion_text",
    ]
    _append_csv_row(INSERTIONS_CSV, fieldnames, row)


def _append_verification(row: Dict[str, Any]) -> None:
    fieldnames = [
        "row_id",
        "timestamp",
        "batch_id",
        "affiliation_id",
        "verified",
        "verification_outcome",
        "remarks",
        "chosen_any",
        "affiliation_has_results",
        "query_has_results",
        "verifier",
    ]
    _append_csv_row(_verification_csv_path(), fieldnames, row)


def _record_to_row(
    batch_id: str,
    record: Dict[str, Any],
    verifications: Dict[str, Dict[str, Any]],
    insertions: Dict[str, str],
    remarks: Dict[str, str],
) -> Optional[Dict[str, Any]]:
    affiliation_id = record.get("affiliation_id") or ""
    row_id = f"{batch_id}::{affiliation_id}"

    verification = verifications.get(row_id, {})
    if verification.get("verified", "false") == "true":
        return None

    affiliation_response = record.get("affiliation_search") or {}
    query_response = record.get("query_search") or {}
    affiliation_items = _extract_items(affiliation_response)
    query_items = _extract_items(query_response)
    top_score, top_substring = _first_item_score_and_substring(affiliation_response)

    return {
        "row_id": row_id,
        "batch_id": batch_id,
        "affiliation_id": affiliation_id,
        "affiliation_string": record.get("affiliation_string", ""),
        "affiliation_count": len(affiliation_items),
        "query_count": len(query_items),
        "score": top_score,
        "substring": top_substring,
        "matching_type": _first_item_matching_type(affiliation_response),
        "chosen_any": _chosen_any(affiliation_response),
        "verified": False,
        "remarks": remarks.get(row_id, ""),
        "insertion_note": insertions.get(row_id, ""),
    }


def _build_rows(batch_id: Optional[str] = None) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    verifications = _load_verifications()
    insertions = _load_insertions_latest()
    remarks = _load_remarks_latest()

    if batch_id:
        batch_dirs = [BATCHES_ROOT / batch_id] if (BATCHES_ROOT / batch_id).is_dir() else []
    else:
        batch_dirs = sorted(p for p in BATCHES_ROOT.iterdir() if p.is_dir()) if BATCHES_ROOT.exists() else []

    for batch_dir in batch_dirs:
        current_batch_id = batch_dir.name
        for output_path in sorted(batch_dir.glob("*.json")):
            record = _json_load(output_path)
            row = _record_to_row(current_batch_id, record, verifications, insertions, remarks)
            if row:
                rows.append(row)

    return rows


def _format_process_error(exc: Exception) -> str:
    message = str(exc)
    lowered = message.lower()
    if (
        "nameresolutionerror" in lowered
        or "failed to resolve" in lowered
        or "name or service not known" in lowered
        or "temporary failure in name resolution" in lowered
    ):
        return (
            "Cannot reach api.ror.org from the container (DNS/network). "
            "Restart with DNS_SERVER_1 set to your server DNS, or uncomment "
            "network_mode: host in docker-compose.yml, and ensure outbound HTTPS "
            "to api.ror.org is allowed."
        )
    return f"ROR processing failed: {exc}"


def _parse_process_affiliations() -> tuple[List[str], Optional[str]]:
    affiliations: List[str] = []
    error: Optional[str] = None

    if request.content_type and "multipart/form-data" in request.content_type:
        upload = request.files.get("file")
        pasted = (request.form.get("text") or "").strip()
        if upload and upload.filename:
            content = upload.read().decode("utf-8-sig")
            affiliations = ror_implementor.parse_affiliation_upload(upload.filename, content)
        elif pasted:
            affiliations = ror_implementor.parse_affiliation_lines(pasted)
        else:
            error = "Paste affiliations or choose a .txt/.csv file."
        return affiliations, error

    payload = request.get_json(silent=True) or {}
    pasted = str(payload.get("text", "")).strip()
    if pasted:
        affiliations = ror_implementor.parse_affiliation_lines(pasted)
    elif isinstance(payload.get("affiliations"), list):
        affiliations = [str(item).strip() for item in payload["affiliations"] if str(item).strip()]
    else:
        error = "Provide affiliations as text (one per line) or upload a .txt/.csv file."

    return affiliations, error


@app.get("/")
def index() -> str:
    return render_template("index.html")


@app.get("/terms")
def terms_of_service() -> str:
    return render_template("terms.html")


@app.get("/security")
def security_compliance() -> str:
    return render_template("security.html")


@app.get("/accessibility")
def accessibility_statement() -> str:
    return render_template("accessibility.html")


@app.post("/api/process")
def api_process() -> Any:
    affiliations, parse_error = _parse_process_affiliations()
    if parse_error:
        return jsonify({"error": parse_error}), 400
    if not affiliations:
        return jsonify({"error": "No affiliations found in the input."}), 400

    batch_id = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H-%M-%SZ")
    output_dir = BATCHES_ROOT / batch_id
    try:
        written = ror_implementor.process_affiliation_strings(affiliations, output_dir)
    except Exception as exc:
        return jsonify({"error": _format_process_error(exc)}), 502

    rows = _build_rows(batch_id)
    return jsonify(
        {
            "ok": True,
            "batch_id": batch_id,
            "processed": len(written),
            "rows": rows,
        }
    )


@app.get("/api/results")
def api_results() -> Any:
    batch_id = request.args.get("batch_id", "").strip() or None
    return jsonify({"rows": _build_rows(batch_id)})


@app.get("/api/result")
def api_result_detail() -> Any:
    batch_id = request.args.get("batch_id", "").strip()
    affiliation_id = request.args.get("affiliation_id", "").strip()
    if not batch_id or not affiliation_id:
        return jsonify({"error": "batch_id and affiliation_id are required"}), 400

    output_path = _result_detail_path(batch_id, affiliation_id)
    if not output_path.exists():
        return jsonify({"error": "result not found"}), 404

    return jsonify(_json_load(output_path))


def _verification_row(
    row_id: str,
    batch_id: str,
    affiliation_id: str,
    record: Dict[str, Any],
    verification_outcome: str,
    remarks: str,
    verifier: str,
) -> Dict[str, Any]:
    affiliation_response = record.get("affiliation_search") or {}
    query_response = record.get("query_search") or {}
    return {
        "row_id": row_id,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "batch_id": batch_id,
        "affiliation_id": affiliation_id,
        "verified": "true",
        "verification_outcome": verification_outcome,
        "remarks": remarks,
        "chosen_any": "true" if _chosen_any(affiliation_response) else "false",
        "affiliation_has_results": "true" if _extract_items(affiliation_response) else "false",
        "query_has_results": "true" if _extract_items(query_response) else "false",
        "verifier": verifier,
    }


@app.post("/api/remarks")
def api_remarks() -> Any:
    payload = request.get_json(silent=True) or {}
    row_id = str(payload.get("row_id", "")).strip()
    remarks = str(payload.get("remarks", "")).strip()

    if "::" not in row_id:
        return jsonify({"error": "invalid row_id"}), 400

    batch_id, affiliation_id = row_id.split("::", 1)
    detail_path = _result_detail_path(batch_id, affiliation_id)
    if not detail_path.exists():
        return jsonify({"error": "result not found"}), 404

    _append_remarks(
        {
            "row_id": row_id,
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "batch_id": batch_id,
            "affiliation_id": affiliation_id,
            "remarks": remarks,
        }
    )
    return jsonify({"ok": True})


@app.post("/api/verify")
def api_verify() -> Any:
    payload = request.get_json(silent=True) or {}
    row_id = payload.get("row_id", "")
    verification_outcome = str(payload.get("verification_outcome", "")).strip()
    remarks = str(payload.get("remarks", "")).strip()
    verifier = str(payload.get("verifier", "")).strip()

    if "::" not in row_id:
        return jsonify({"error": "invalid row_id"}), 400

    if verification_outcome not in ("verified", "no_result_found"):
        return jsonify({"error": "verification_outcome must be verified or no_result_found"}), 400

    batch_id, affiliation_id = row_id.split("::", 1)
    detail_path = _result_detail_path(batch_id, affiliation_id)
    if not detail_path.exists():
        return jsonify({"error": "result not found"}), 404

    record = _json_load(detail_path)
    if remarks:
        _append_remarks(
            {
                "row_id": row_id,
                "timestamp": datetime.now(timezone.utc).isoformat(),
                "batch_id": batch_id,
                "affiliation_id": affiliation_id,
                "remarks": remarks,
            }
        )
    _append_verification(
        _verification_row(
            row_id, batch_id, affiliation_id, record, verification_outcome, remarks, verifier
        )
    )
    return jsonify({"ok": True})


@app.post("/api/verify/bulk")
def api_verify_bulk() -> Any:
    payload = request.get_json(silent=True) or {}
    row_ids = payload.get("row_ids") or []
    verification_outcome = str(payload.get("verification_outcome", "verified")).strip()
    verifier = str(payload.get("verifier", "")).strip()
    if not isinstance(row_ids, list):
        return jsonify({"error": "row_ids must be a list"}), 400
    if verification_outcome not in ("verified", "no_result_found"):
        return jsonify({"error": "verification_outcome must be verified or no_result_found"}), 400

    verified_count = 0
    for row_id in row_ids:
        if not isinstance(row_id, str) or "::" not in row_id:
            continue
        batch_id, affiliation_id = row_id.split("::", 1)
        detail_path = _result_detail_path(batch_id, affiliation_id)
        if not detail_path.exists():
            continue

        record = _json_load(detail_path)
        _append_verification(
            _verification_row(
                row_id, batch_id, affiliation_id, record, verification_outcome, "", verifier
            )
        )
        verified_count += 1

    return jsonify({"ok": True, "verified_count": verified_count})


@app.post("/api/insertion")
def api_insertion() -> Any:
    payload = request.get_json(silent=True) or {}
    row_id = str(payload.get("row_id", "")).strip()
    insertion_record = payload.get("insertion_record")
    insertion_text = str(payload.get("insertion_text", "")).strip()

    if "::" not in row_id:
        return jsonify({"error": "invalid row_id"}), 400

    if isinstance(insertion_record, dict):
        normalized, err = _normalize_insertion_record(insertion_record)
        if err:
            return jsonify({"error": err}), 400
        insertion_text = json.dumps(normalized, ensure_ascii=False)
    elif not insertion_text:
        return jsonify({"error": "insertion_record or insertion_text is required"}), 400

    batch_id, affiliation_id = row_id.split("::", 1)
    detail_path = _result_detail_path(batch_id, affiliation_id)
    if not detail_path.exists():
        return jsonify({"error": "result not found"}), 404

    row = {
        "row_id": row_id,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "batch_id": batch_id,
        "affiliation_id": affiliation_id,
        "insertion_text": insertion_text,
    }
    _append_insertion(row)
    return jsonify({"ok": True})


def _load_remarks_export_rows() -> List[Dict[str, str]]:
    if not REMARKS_CSV.exists():
        return []
    latest: Dict[str, Dict[str, str]] = {}
    with REMARKS_CSV.open("r", newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            if _is_csv_header_data_row(row):
                continue
            rid = row.get("row_id") or ""
            if rid:
                latest[rid] = row

    export_rows: List[Dict[str, str]] = []
    for row_id, row in latest.items():
        batch_id = row.get("batch_id") or ""
        affiliation_id = row.get("affiliation_id") or ""
        if "::" in row_id and (not batch_id or not affiliation_id):
            batch_id, affiliation_id = row_id.split("::", 1)

        affiliation_string = ""
        if batch_id and affiliation_id:
            detail_path = _result_detail_path(batch_id, affiliation_id)
            if detail_path.exists():
                record = _json_load(detail_path)
                affiliation_string = str(record.get("affiliation_string") or "")

        export_rows.append(
            {
                "row_id": row_id,
                "batch_id": batch_id,
                "affiliation_id": affiliation_id,
                "affiliation_string": affiliation_string,
                "remarks": (row.get("remarks") or "").strip(),
                "timestamp": row.get("timestamp") or "",
            }
        )
    return export_rows


def _build_remarks_csv() -> io.BytesIO:
    rows = _load_remarks_export_rows()
    text_buffer = io.StringIO()
    fieldnames = [
        "row_id",
        "batch_id",
        "affiliation_id",
        "affiliation_string",
        "remarks",
        "timestamp",
    ]
    writer = csv.DictWriter(text_buffer, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
    buffer = io.BytesIO(text_buffer.getvalue().encode("utf-8"))
    buffer.seek(0)
    return buffer


@app.get("/api/insertions/download")
def api_insertions_download() -> Any:
    records = _load_insertion_records_list()
    if not records:
        return jsonify({"error": "No insertion records saved yet"}), 404
    try:
        buffer = _build_insertions_xlsx(records)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 500
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=INSERTIONS_DOWNLOAD_NAME,
    )


@app.get("/api/remarks/download")
def api_remarks_download() -> Any:
    if not _load_remarks_export_rows():
        return jsonify({"error": "No remarks saved yet"}), 404
    buffer = _build_remarks_csv()
    return send_file(
        buffer,
        mimetype="text/csv",
        as_attachment=True,
        download_name="ror_remarks.csv",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=True)
